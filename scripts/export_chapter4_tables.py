from __future__ import annotations

from pathlib import Path

import pandas as pd

from english_labels import LEVEL1_EN, LEVEL2_EN, TYPE_EN, indicator_label, load_country_names, source_label


PROJECT_DIR = Path(__file__).resolve().parents[1]
YEARS = list(range(2001, 2022))


def blank_repeats(df: pd.DataFrame, cols: list[str]) -> pd.DataFrame:
    """Blank repeated hierarchical labels for Word-friendly table copies."""
    out = df.copy()
    for col in cols:
        out.loc[out[col].eq(out[col].shift()), col] = ""
    return out


def build_table_4_2() -> pd.DataFrame:
    variables = pd.read_excel(PROJECT_DIR / "Variables Chosen.xlsx", sheet_name="Index")
    source = (
        variables.query("变量类型 == '指标体系'")
        [["一级指标", "二级指标", "三级指标", "Variables", "来源"]]
        .copy()
        .reset_index(drop=True)
    )
    table = pd.DataFrame(
        {
            "First-level dimension": source["一级指标"].map(LEVEL1_EN),
            "Second-level dimension": source["二级指标"].map(LEVEL2_EN),
            "Third-level indicator": source.apply(indicator_label, axis=1),
            "Source": source["来源"].map(source_label),
        }
    )
    return blank_repeats(table, ["First-level dimension", "Second-level dimension"])


def build_table_4_3() -> pd.DataFrame:
    selected = pd.read_excel(PROJECT_DIR / "selected_indicator_system_variables.xlsx")
    weights = pd.read_csv(PROJECT_DIR / "critic_weight_two_stage.csv", encoding="utf-8-sig")

    table = selected.merge(
        weights[["Variables", "two_stage_global_weight"]],
        on="Variables",
        how="left",
        validate="one_to_one",
    )
    if table["two_stage_global_weight"].isna().any():
        missing = table.loc[table["two_stage_global_weight"].isna(), "Variables"].tolist()
        raise ValueError(f"Missing two-stage weights for: {missing}")

    level1_labels: dict[str, str] = {}
    level2_labels: dict[str, str] = {}
    next_level1 = 1
    next_level2 = 1
    rows = []

    for idx, row in table.reset_index(drop=True).iterrows():
        level1 = row["一级指标"]
        level2 = row["二级指标"]

        if level1 not in level1_labels:
            level1_labels[level1] = f"A{next_level1} {LEVEL1_EN[level1]}"
            next_level1 += 1
        level2_key = f"{level1}|{level2}"
        if level2_key not in level2_labels:
            level2_labels[level2_key] = f"B{next_level2} {LEVEL2_EN[level2]}"
            next_level2 += 1

        rows.append(
            {
                "First-level dimension": level1_labels[level1],
                "Second-level dimension": level2_labels[level2_key],
                "Third-level indicator": f"C{idx + 1} {indicator_label(row)}",
                "Direction": TYPE_EN[row["类型"]],
                "Weight": round(float(row["two_stage_global_weight"]), 4),
            }
        )

    return blank_repeats(pd.DataFrame(rows), ["First-level dimension", "Second-level dimension"])


def build_table_4_4() -> pd.DataFrame:
    index_data = pd.read_csv(PROJECT_DIR / "index_data_two_stage_comparison.csv", encoding="utf-8-sig")
    index_data = index_data[index_data["Year"].isin(YEARS)].copy()
    index_data = index_data.merge(load_country_names(PROJECT_DIR), on="Alpha-3 code", how="left")
    index_data["SDI_score"] = index_data["SDI_TwoStage"] * 100

    table = (
        index_data.pivot_table(
            index="Country",
            columns="Year",
            values="SDI_score",
            aggfunc="first",
        )
        .reindex(columns=YEARS)
        .round(1)
    )

    sort_year = YEARS[-1]
    table = table.sort_values(sort_year, ascending=False)
    table = table.reset_index()
    table.columns = [str(c) if isinstance(c, int) else c for c in table.columns]
    return table


def autosize_workbook(path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font

    wb = load_workbook(path)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="center", wrap_text=True)

        for column_cells in ws.columns:
            column_letter = column_cells[0].column_letter
            max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
            ws.column_dimensions[column_letter].width = min(max(max_len + 2, 10), 32)

        if ws.title.startswith("Table 4.4"):
            ws.column_dimensions["A"].width = 16
            for col in range(2, ws.max_column + 1):
                ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 8

    wb.save(path)


def main() -> None:
    table_4_2 = build_table_4_2()
    table_4_3 = build_table_4_3()
    table_4_4 = build_table_4_4()

    output_xlsx = PROJECT_DIR / "chapter4_tables.xlsx"
    with pd.ExcelWriter(output_xlsx, engine="openpyxl") as writer:
        table_4_2.to_excel(writer, sheet_name="Table 4.2 Candidate indicators", index=False)
        table_4_3.to_excel(writer, sheet_name="Table 4.3 Final weights", index=False)
        table_4_4.to_excel(writer, sheet_name="Table 4.4 SDI 2001-2021", index=False)
        table_4_4[["Country"] + [str(y) for y in YEARS[:11]]].to_excel(
            writer,
            sheet_name="Table 4.4 2001-2011",
            index=False,
        )
        table_4_4[["Country"] + [str(y) for y in YEARS[11:]]].to_excel(
            writer,
            sheet_name="Table 4.4 2012-2021",
            index=False,
        )

    autosize_workbook(output_xlsx)

    csv_outputs = {
        "chapter4_table_4_2.csv": table_4_2,
        "chapter4_table_4_3.csv": table_4_3,
        "chapter4_table_4_4.csv": table_4_4,
    }
    for filename, table in csv_outputs.items():
        table.to_csv(PROJECT_DIR / filename, index=False, encoding="utf-8-sig")

    print(f"Exported {output_xlsx}")
    print(f"Table 4.2 rows: {len(table_4_2)}")
    print(f"Table 4.3 rows: {len(table_4_3)}")
    print(f"Table 4.4 countries: {len(table_4_4)}, years: {YEARS[0]}-{YEARS[-1]}")


if __name__ == "__main__":
    main()
